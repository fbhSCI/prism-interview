from pathlib import Path
import json
import numpy as np
import pytest

from zoom_copilot.answers import parse_reply, looks_like_question
from zoom_copilot.audio import Segmenter, RATE
from zoom_copilot.knowledge import KnowledgeBase, Passage, chunks, tokenize


def test_question_detection():
    assert looks_like_question("How do you handle calibration uncertainty?")
    assert looks_like_question("Could you describe the spectral sensing operator")
    assert not looks_like_question("We should move to the next slide.")
    assert not looks_like_question("Okay")
    assert looks_like_question("I would like to understand how your research connects to inverse rendering.")
    assert looks_like_question("You mentioned calibration. Please discuss its uncertainty.")
    assert looks_like_question("You mentioned calibration. How would you validate it.")
    assert not looks_like_question("I would like to move to the next slide.")


def test_knowledge_retrieval_local(tmp_path):
    kb = KnowledgeBase(tmp_path / "index.json")
    one = tmp_path / "stokes.txt"
    two = tmp_path / "materials.txt"
    one.write_text("Stokes inversion uses a polarization calibration matrix and analyzer-channel intensities."
                   " A sensing operator first reconstructs spectral images, then Stokes parameters are retrieved.", encoding="utf-8")
    two.write_text("BRDF is related to material reflectance under varying illumination and viewing angles."
                   " Calibration uncertainty affects future material inference research.", encoding="utf-8")
    count, errors = kb.add_files([str(one), str(two)])
    assert count == 2
    assert not errors
    results = kb.search("Stokes polarization calibration matrix", top_k=1)
    assert results and Path(results[0].source).name == "stokes.txt"
    reload = KnowledgeBase(tmp_path / "index.json")
    assert len(reload.document_names) == 2


def test_parse_ollama_json():
    raw = '```json\n{"question_zh":"为什么？","answer_en":"Because calibration matters.","answer_zh":"因为标定很重要。"}\n```'
    res = parse_reply(raw, [Passage("C:/test/paper.pdf", "第 1 页", "hello")])
    assert "calibration" in res.answer_en
    # A retrieved candidate is not an actual citation unless the model cites it.
    assert res.sources == []
    with pytest.raises(ValueError):
        parse_reply('{"question_zh":"ok"}', [])


def test_segmenter_finalizes_utterance():
    # A pure sine is not human speech. This test isolates segmentation while
    # separate real-VAD tests verify that production rejects non-speech audio.
    seg = Segmenter(threshold=0.01, silence_sec=0.4, vad=False)
    events = []
    # A loud voiced segment then a quiet trailing pause.
    speech = (np.sin(2 * np.pi * 220 * np.arange(RATE) / RATE) * 0.1).astype(np.float32)
    for block in np.array_split(speech, 5):
        events.extend(seg.push(block))
    for _ in range(4):
        events.extend(seg.push(np.zeros(RATE // 5, dtype=np.float32)))
    assert any(kind == "final" for kind, _ in events)
    assert seg.flush() is None


def test_chunking():
    parts = list(chunks("abc " * 400, maximum=160, overlap=30))
    assert len(parts) > 2
    assert all(p for p in parts)


def test_pdf_docx_pptx_extraction(tmp_path):
    from zoom_copilot.knowledge import read_document
    from public_fixtures import write_synthetic_pdf
    from docx import Document
    from pptx import Presentation

    pdf = tmp_path / "research.pdf"
    write_synthetic_pdf(pdf, [(72, 720, "Project analysis and validation evidence.")])

    word = tmp_path / "research.docx"
    d = Document()
    d.add_paragraph("Spectral polarimetric reconstruction with Stokes calibration.")
    d.save(str(word))

    slide = tmp_path / "research.pptx"
    pr = Presentation()
    sl = pr.slides.add_slide(pr.slide_layouts[6])
    from pptx.util import Inches
    box = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    box.text = "Material inference and BRDF uncertainty."
    pr.save(str(slide))

    assert "Project analysis" in read_document(pdf)[0][1]
    assert "Stokes calibration" in read_document(word)[0][1]
    assert "BRDF" in read_document(slide)[0][1]


def test_ollama_chat_request_shape(monkeypatch):
    from zoom_copilot import answers
    from urllib.parse import urlparse

    class FakeResponse:
        def __init__(self, body):
            self.body = body
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False
        def read(self, *_):
            return self.body

    def fake_urlopen(request, timeout=0):
        assert urlparse(request.full_url).hostname == "127.0.0.1"
        payload = json.loads(request.data.decode("utf-8"))
        assert payload["model"] == "qwen2.5:3b"
        assert payload["format"] == "json"
        assert "Stokes" in payload["messages"][1]["content"]
        reply = {"message": {"content": json.dumps({
            "question_zh": "解释Stokes", "answer_en": "Stokes is distinct from reconstruction.",
            "answer_zh": "Stokes 与重建不同。"}, ensure_ascii=False)}}
        return FakeResponse(json.dumps(reply).encode())

    monkeypatch.setattr(answers, "urlopen", fake_urlopen)
    answer = answers.OllamaClient().answer("Explain Stokes inversion", [])
    assert "Stokes" in answer.answer_en
    assert "Stokes" in answer.answer_zh
