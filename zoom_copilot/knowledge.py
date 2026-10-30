"""Local, page-preserving research retrieval.

The default is cached BM25 with explicit English/Chinese terminology expansion,
not semantic/vector retrieval. An application may explicitly attach a *local*
embedding callable with ``set_embedder`` to enable reciprocal-rank hybrid search.
This module never downloads models or sends documents over the network.
Returned passages are candidate evidence, not a claim that an answer cited them.
"""
from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass, replace
import hashlib
import json
import math
import os
from pathlib import Path
import re
import tempfile
import threading
import unicodedata
from typing import Callable, Iterable, Sequence


VALID_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}
_EXTRACTION_PIPELINE = "pypdf-page-order-v6-docx-sections"
_WORDS = re.compile(r"[a-z]+[a-z0-9]*|\d+(?:\.\d+)?|[\u4e00-\u9fff]+", re.I)
_STOP = set("""a an the and or but of to for in on at by with from into as is are
was were be been being it its this that these those i my me we our you your he
she they their them can could would should do does did have has had will may
might must not no than then there here how why what which who when where
please tell explain describe discuss compare about between also more most
some any each all such so if both through using use used based paper papers
research study work proposed approach method methods system results result
question answer give including include includes respectively specifically
""".split())
# Canonical spellings, not a general-purpose stemmer. These are bidirectional.
_CANON = {
    "polarimetric": "polarization", "polarimetry": "polarization",
    "polarisation": "polarization", "polarizations": "polarization",
    "spectra": "spectral", "spectrum": "spectral", "spectrally": "spectral",
    "hyperspectral": "hyperspectral", "analyser": "analyzer",
    "analysers": "analyzer", "analyzers": "analyzer", "matrices": "matrix",
    "calibrated": "calibration", "calibrating": "calibration",
    "reconstruct": "reconstruction", "reconstructed": "reconstruction",
    "reconstructing": "reconstruction", "measurements": "measurement",
    "measured": "measurement", "measuring": "measurement",
    "materials": "material", "bands": "band", "channels": "channel",
    "fields": "field", "normals": "normal", "geometry": "geometry",
    "geometric": "geometry", "scattering": "scattering", "scatter": "scattering",
    "wavelengths": "wavelength", "operators": "operator", "responses": "response",
    "inversion": "inverse", "rendered": "rendering", "render": "rendering",
    "snapshots": "snapshot", "constraints": "constraint", "ghostimaging": "ghost",
    "reflectances": "reflectance", "brdfs": "brdf", "stokesvector": "stokes",
    "unchanged": "fixed", "recalibration": "calibration", "recalibrating": "calibration",
    "conditions": "condition", "speeds": "speed", "limitations": "limitation",
    "demodulate": "demodulation", "demodulated": "demodulation",
    "demodulator": "demodulation", "demodulating": "demodulation",
}
_ZH_TERMS = {
    "项目管理": "project management", "团队合作": "teamwork collaboration",
    "工作经历": "work experience", "职业目标": "career goals",
    "优势": "strength skills", "缺点": "weakness improvement",
    "简历": "resume curriculum vitae", "岗位": "role position job",
    "供应链": "supply chain", "现金流": "cash flow",
    "客户体验": "customer experience", "临床试验": "clinical trial",
    "知情同意": "informed consent", "置信区间": "confidence interval",
    "统计显著性": "statistical significance", "根因分析": "root cause analysis",
    "计算成像": "computational imaging", "鬼成像": "ghost gisc",
    "关联成像": "ghost gisc", "光谱": "spectral", "高光谱": "hyperspectral spectral",
    "偏振": "polarization", "斯托克斯": "stokes", "斯托克": "stokes",
    "感知算子": "sensing operator", "测量算子": "sensing operator",
    "传感算子": "sensing operator", "传感矩阵": "sensing matrix",
    "仪器矩阵": "instrument matrix", "标定矩阵": "calibration instrument matrix",
    "标定": "calibration", "水下": "underwater", "散射": "scattering",
    "探测器域": "detector domain", "图像域": "image domain", "退化": "degradation",
    "气流": "airflow", "无风": "noflow", "无气流": "noflow",
    "固定": "fixed unchanged", "衰减": "attenuation", "背景光": "background",
    "逆渲染": "inverse rendering", "材质": "material", "材料": "material",
    "反射": "reflectance", "几何": "geometry", "法线": "normal",
    "粗糙度": "roughness", "金属度": "metallicity", "反照率": "albedo",
    "双向反射分布": "brdf", "波段": "band wavelength", "波长": "wavelength",
    "采样": "sampling", "重建": "reconstruction", "神经场": "neural field",
    "快照": "snapshot", "单次曝光": "single shot", "单拍": "single shot",
    "全穆勒矩阵": "full mueller matrix", "穆勒": "mueller",
    "不确定性": "uncertainty error", "局限": "limitation", "限制": "limitation",
    "三域": "three domain frequency spectral spatial", "频域": "frequency domain",
    "空间域": "spatial domain", "频谱域": "spectral domain", "光谱域": "spectral domain",
}
_ALIASES = {
    "gisc": ("ghost", "sparsity"), "ghost": ("gisc",),
    "hpbrdf": ("hyperspectral", "brdf", "mueller"),
    "pbrdf": ("brdf", "polarization"), "brdf": ("reflectance",),
    "dolp": ("linear", "polarization"), "aolp": ("angle", "polarization"),
}

# These words can rank useful passages but are not sufficient evidence that a
# passage addresses a named subject. This prevents an unrelated algorithm from
# retrieving an imaging paper only because both discuss "algorithm complexity".
_GENERIC_ANCHORS = set("""algorithm algorithms complexity model models data dataset
evaluation evaluate evaluated performance accuracy accurate experiment experiments
measurement measured parameter parameters error errors number integer bit computational
develop developed main key contribution contributions personally already published
publication publications three two first phd doctoral master experience background
research context local document lookup scientific concept theoretical theory analysis
full only recover recovery retrieval including observe observed mean improvement
gain universal average numerical using method approach proposed support relate
future completed existing own relevant compare comparison question explain answer
""".split())
_CONTEXT_MARKER = '\nCandidate research context for local document lookup:\n'


def _remove_contact_details(text: str) -> str:
    """Omit irrelevant CV contacts from selected evidence, preserving raw index.

    Do not match arbitrary long digit strings: instrument values, equations and
    scientific quantities must survive. Phone stripping requires contact context
    or an international '+' prefix and at least ten digits.
    """
    email = re.compile(r'\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b', re.I)
    phone = re.compile(r'(?<![\w.])\+?\d[\d ()-]{8,}\d(?![\w.])')
    lines=[]
    for line in text.splitlines():
        contact = bool(email.search(line) or re.search(r'\b(?:tel|phone|mobile)\s*[:：]|(?:电话|手机)\s*[:：]',line,re.I))
        line=email.sub('',line)
        line=phone.sub(lambda m:'' if len(re.sub(r'\D','',m[0]))>=10 and
                       (contact or (m[0].startswith('+') and len(re.findall(r'[ -]+',m[0]))>=2))
                       else m[0],line)
        line=re.sub(r'\s*\|\s*(?=\||$)','',line).strip()
        if line and not re.fullmatch(r'(?:Google Scholar|LinkedIn)(?:\s*\|\s*(?:Google Scholar|LinkedIn))*',line,re.I):
            lines.append(line)
    return '\n'.join(lines)


@dataclass
class Passage:
    source: str
    location: str
    text: str
    passage_id: str = ""
    score: float = 0.0


def _normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).lower().replace("\x00", "")
    # PDF column wrapping often hyphenates a word across a line.
    return re.sub(r"(?<=[a-z])[-\u00ad]\s*\n\s*(?=[a-z])", "", text)


def tokenize(text: str) -> list[str]:
    out = []
    normalized = _normalize(text)
    # Keep the component words AND the compound. In scientific text, "no-flow"
    # and "single-shot" have a meaning lost by deleting the short prefix.
    out.extend(a + b for a, b in re.findall(r"\b([a-z]{2,})[-–]([a-z]{2,})\b", normalized))
    for word in _WORDS.findall(normalized):
        if word in _STOP:
            continue
        if "\u4e00" <= word[0] <= "\u9fff":
            # Bigrams preserve Chinese phrases without a resident tokenizer model.
            out.extend(word[i:i + 2] for i in range(len(word) - 1))
            if len(word) == 1:
                out.append(word)
        else:
            out.append(_CANON.get(word, word))
    return out


def _query_weights(query: str) -> dict[str, float]:
    weights = dict.fromkeys(tokenize(query), 1.0)
    normalized = _normalize(query)
    for phrase, expansion in _ZH_TERMS.items():
        if phrase in normalized:
            for term in tokenize(expansion):
                weights[term] = max(weights.get(term, 0.0), 0.85)
    for term in list(weights):
        for alias in _ALIASES.get(term, ()):
            weights[alias] = max(weights.get(alias, 0.0), 0.3)
    # "Stokes calibration matrix" can mean the forward instrument response or
    # its inverse/demodulation. Recall both definitions rather than rank only
    # generic scene-encoding passages. This is terminology, never a file prior.
    if "stokes" in weights and ("calibration" in weights or "matrix" in weights):
        for term, weight in (("instrument", 0.9), ("demodulation", 0.9), ("pseudoinverse", 0.9)):
            weights[term] = max(weights.get(term, 0.0), weight)
    return weights


def read_document(path: Path) -> list[tuple[str, str]]:
    """Read text only. Page numbers are PDF file pages (1-based), not printed folios.

    Images/formula OCR is deliberately not inferred. Empty image-only PDFs report
    an import issue. Native PDF readers are imported lazily for testability.
    """
    suffix = path.suffix.lower()
    if suffix not in VALID_EXTENSIONS:
        raise ValueError(f"不支持的文件格式：{suffix}")
    if path.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("单个文件不能超过 100 MB；请先提取相关章节。")
    if suffix in (".txt", ".md"):
        return [("全文", path.read_text(encoding="utf-8-sig", errors="replace"))]
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF 已加密，请先提供可读取的副本。")
        if len(reader.pages) > 500:
            raise ValueError("PDF 超过 500 页，请只添加相关章节。")
        # Plain mode follows native content-stream order. Layout extraction can
        # interleave columns; scanned PDFs still require user-provided OCR text.
        return [(f"第 {i + 1} 页", page.extract_text(extraction_mode="plain") or "")
                for i, page in enumerate(reader.pages)]
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Seminar scripts often carry literal Slide 01 / Q1 headings. They are
        # useful, truthful locations; never call them DOCX page numbers. Keep
        # each section separate so a result/answer cannot borrow the next slide.
        section_heading = re.compile(r"^(?:Slide\s+\d{1,3}\b|Q\d{1,2}[.、:：])", re.I)
        if sum(bool(section_heading.match(part.strip())) for part in parts) >= 2:
            sections, body, label = [], [], "开头（无固定页码）"
            for part in parts:
                if section_heading.match(part.strip()):
                    if body:
                        sections.append((label, "\n".join(body)))
                    label, body = part.strip()[:150], [part]
                else:
                    body.append(part)
            if body:
                sections.append((label, "\n".join(body)))
            for index, table in enumerate(doc.tables, 1):
                sections.append((f"表 {index}（无固定页码）", "\n".join(
                    " | ".join(c.text for c in row.cells) for row in table.rows)))
            return sections
        for table in doc.tables:
            parts.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
        return [("全文（无固定页码）", "\n".join(parts))]
    if suffix == ".pptx":
        from pptx import Presentation
        result = []
        for index, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
            result.append((f"Slide {index}", "\n".join(texts)))
            if slide.has_notes_slide:
                note = slide.notes_slide.notes_text_frame.text
                if note.strip():
                    result.append((f"Slide {index} 讲者备注", note))
        return result
    raise ValueError("不支持的文件格式")


def chunks(text: str, maximum: int = 1100, overlap: int = 180) -> Iterable[str]:
    """Overlapping chunks remain within their original page/slide boundary."""
    if maximum < 40 or overlap < 0 or overlap >= maximum:
        raise ValueError("分段长度需至少 40，且 0 <= overlap < maximum。")
    text = re.sub(r"(?<=[A-Za-z])[-\u00ad]\s*\n\s*(?=[a-z])", "", text)
    clean = re.sub(r"[ \t]+", " ", text.replace("\x00", ""))
    clean = re.sub(r"\n{3,}", "\n\n", clean).strip()
    start = 0
    while start < len(clean):
        end = min(start + maximum, len(clean))
        if end < len(clean):
            lower = start + maximum // 2
            boundaries = [clean.rfind(mark, lower, end) for mark in ("\n\n", ". ", "。", "\n", "; ")]
            boundary = max(boundaries)
            if boundary >= lower:
                end = boundary + 1
        piece = clean[start:end].strip()
        if piece:
            yield piece
        if end >= len(clean):
            break
        start = max(end - overlap, start + 1)


def _passage(source: str, location: str, text: str) -> Passage:
    digest = hashlib.sha256((source + "\0" + location + "\0" + text).encode("utf-8")).hexdigest()[:20]
    return Passage(source, location, text, digest)


class KnowledgeBase:
    """Thread-safe, atomically saved local index with content-hash refresh.

    Import/clear operations serialize. Search reads a consistent old/new snapshot
    while documents are parsed. A clear submitted during import runs afterwards,
    so an in-flight import cannot resurrect the cleared knowledge base.
    """
    def __init__(self, index_path: Path):
        self.index_path = Path(index_path)
        self._lock = threading.RLock()
        self._write_lock = threading.RLock()
        self.passages: list[Passage] = []
        self.document_names: list[str] = []
        self.document_records: dict[str, dict] = {}
        self.load_warning = ""
        self.embedding_warning = ""
        self._embedder = None
        self._embedding_name = ""
        self._vectors: list[list[float]] = []
        self._revision = 0
        self._load()
        self._rebuild_index()

    @property
    def retrieval_mode(self) -> str:
        return "hybrid-local" if self._embedder is not None and self._vectors else "bm25-local-terminology"

    def _load(self) -> None:
        if not self.index_path.is_file():
            return
        try:
            data = json.loads(self.index_path.read_text(encoding="utf-8"))
            passages = [_passage(str(p["source"]), str(p["location"]), str(p["text"]))
                        for p in data.get("passages", [])]
            if len(passages) > 20000:
                raise ValueError("索引超过 20000 段上限")
            self.passages = passages
            self.document_names = list(dict.fromkeys(str(x) for x in data.get("documents", [])))
            self.document_records = data.get("document_records", {})
        except (ValueError, TypeError, KeyError, OSError) as ex:
            self.load_warning = f"本地索引读取失败，原文件保留：{ex}"

    def _save(self) -> None:
        self.index_path.parent.mkdir(parents=True, exist_ok=True)
        temp_name = ""
        try:
            with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=self.index_path.parent,
                                             prefix=self.index_path.name + ".", suffix=".tmp", delete=False) as out:
                temp_name = out.name
                json.dump({"version": 2, "passages": [asdict(p) for p in self.passages],
                           "documents": self.document_names, "document_records": self.document_records},
                          out, ensure_ascii=False)
                out.flush()
                os.fsync(out.fileno())
            os.replace(temp_name, self.index_path)
        finally:
            if temp_name and Path(temp_name).exists():
                Path(temp_name).unlink()

    def _rebuild_index(self) -> None:
        # Called under _lock after every mutation; tokenization is cached per revision.
        self._counts = tuple(Counter(tokenize(p.text)) for p in self.passages)
        self._lengths = tuple(sum(c.values()) for c in self._counts)
        self._frequency = Counter()
        for count in self._counts:
            self._frequency.update(count.keys())
        self._mean_length = sum(self._lengths) / max(len(self._lengths), 1)
        self._titles = tuple(set(tokenize(Path(p.source).stem)) for p in self.passages)
        self._title_phrases = tuple("".join(c for c in _normalize(Path(p.source).stem) if c.isalnum())
                                    for p in self.passages)
        self._vectors = []
        self._revision += 1

    def _commit(self, passages: list[Passage], names: list[str], records: dict[str, dict]) -> None:
        """Publish only a state that was successfully saved; caller holds _lock."""
        previous = self.passages, self.document_names, self.document_records
        self.passages, self.document_names, self.document_records = passages, names, records
        try:
            self._save()
        except Exception:
            self.passages, self.document_names, self.document_records = previous
            raise
        self._rebuild_index()

    def add_files(self, paths: list[str]) -> tuple[int, list[str]]:
        added, issues = 0, []
        with self._write_lock:
            for raw in paths:
                try:
                    path = Path(raw).expanduser().resolve(strict=True)
                    source = str(path)
                    stat = path.stat()
                    if stat.st_size > 100 * 1024 * 1024:
                        raise ValueError("单个文件不能超过 100 MB。")
                    with path.open("rb") as content:
                        digest = hashlib.file_digest(content, "sha256").hexdigest()
                    with self._lock:
                        old = self.document_records.get(source, {})
                    if old.get("sha256") == digest and old.get("pipeline") == _EXTRACTION_PIPELINE:
                        issues.append(f"内容未变化，已跳过：{path.name}")
                        continue
                    extracted = read_document(path)
                    fresh = [_passage(source, label, piece) for label, text in extracted for piece in chunks(text)]
                    if not fresh or sum(len(p.text.strip()) for p in fresh) < 30:
                        issues.append(f"没有可检索文字：{path.name}（扫描图片需先 OCR，原索引保留）")
                        continue
                    with self._lock:
                        retained = [p for p in self.passages if p.source != source]
                        if len(retained) + len(fresh) > 20000:
                            raise ValueError("超出 20000 个文本段落上限")
                        record = {"sha256": digest, "mtime_ns": stat.st_mtime_ns,
                                  "size": stat.st_size, "passages": len(fresh), "pipeline": _EXTRACTION_PIPELINE}
                        self._commit(retained + fresh, list(dict.fromkeys(self.document_names + [source])),
                                     {**self.document_records, source: record})
                    empty_pages = [label for label, text in extracted if not text.strip()]
                    if empty_pages:
                        issues.append(f"部分页面无可提取文字（未 OCR）：{path.name}：{', '.join(empty_pages[:8])}")
                    added += len(fresh)
                except Exception as ex:
                    issues.append(f"{Path(raw).name}: {ex}")
        if self._embedder is not None:
            self._refresh_vectors()
        return added, issues

    def clear(self) -> None:
        with self._write_lock, self._lock:
            self._commit([], [], {})

    def set_embedder(self, encode: Callable[[list[str]], Sequence[Sequence[float]]] | None,
                     name: str = "explicit-local-model") -> None:
        """Opt-in local embedding interface. Caller must supply an offline encoder.

        No cloud/API encoder is constructed here. Vectors stay in memory. Failures
        fall back to BM25 and are exposed in embedding_warning/retrieval_mode.
        """
        with self._lock:
            self._embedder, self._embedding_name, self._vectors = encode, name, []
        if encode is not None:
            self._refresh_vectors()

    @staticmethod
    def _unit(vector: Sequence[float]) -> list[float]:
        values = [float(x) for x in vector]
        if not values or not all(math.isfinite(x) for x in values):
            raise ValueError("无效本地嵌入向量")
        norm = math.sqrt(sum(x * x for x in values))
        return [x / norm for x in values] if norm else [0.0] * len(values)

    def _refresh_vectors(self) -> None:
        with self._lock:
            encoder, revision = self._embedder, self._revision
            texts = [p.text for p in self.passages]
        if encoder is None or not texts:
            return
        try:
            vectors = [self._unit(v) for v in encoder(texts)]
            if len(vectors) != len(texts) or len({len(v) for v in vectors}) != 1:
                raise ValueError("本地嵌入向量数量/维度不一致")
            with self._lock:
                if revision == self._revision and encoder is self._embedder:
                    self._vectors, self.embedding_warning = vectors, ""
        except Exception as ex:
            with self._lock:
                self._vectors, self.embedding_warning = [], f"嵌入失败，回退本地 BM25：{ex}"

    def search(self, query: str, top_k: int = 3) -> list[Passage]:
        if not query.strip() or top_k <= 0:
            return []
        with self._lock:
            current, counts = tuple(self.passages), self._counts
            lengths, frequency, mean = self._lengths, self._frequency, self._mean_length
            titles, title_phrases, vectors, encoder = self._titles, self._title_phrases, self._vectors, self._embedder
        if not current:
            return []
        primary, separator, expansion = query.partition(_CONTEXT_MARKER)
        primary_weights = _query_weights(primary)
        weights = dict(primary_weights)
        if separator:
            for term, weight in _query_weights(expansion).items():
                weights[term] = max(weights.get(term,0.0), weight * .55)
        query_phrase = "".join(c for c in _normalize(query) if c.isalnum())
        if not weights:
            return []
        n = len(current)
        # Discard common filler; aliases must not be the only weak lexical signal.
        primary_subjects = {w for w, weight in primary_weights.items()
                            if weight >= .8 and w not in _GENERIC_ANCHORS and not w.isdigit()}
        # Distinctive words absent from the library cannot be silently replaced
        # by generic algorithm/data matches. Generic-only questions can still
        # use normal BM25, and local profile expansion may anchor personal asks.
        subject_weights = {w:weight for w,weight in weights.items()
                           if w not in _GENERIC_ANCHORS and not w.isdigit()}
        if primary_subjects and not any(frequency[w] for w in primary_subjects):
            return []
        anchor_weights = subject_weights or weights
        anchors = {w for w, weight in anchor_weights.items() if weight >= (.5 if separator else .8)
                   and 0 < frequency[w] < n * 0.75 + 1}
        scores = []
        for i, count in enumerate(counts):
            matches = anchors.intersection(count)
            if not matches:
                continue
            score = 0.0
            for term, weight in weights.items():
                tf = count[term]
                if tf:
                    df = frequency[term]
                    idf = math.log1p((n - df + 0.5) / (df + 0.5))
                    score += weight * idf * tf * 2.2 / (tf + 1.2 * (0.25 + 0.75 * lengths[i] / max(mean, 1)))
            # Multi-concept questions benefit from covering distinct anchor terms.
            score *= 0.75 + 0.5 * len(matches) / max(len(anchors), 1)
            score += 0.12 * len(anchors.intersection(titles[i]))
            # Explicit full paper titles disambiguate a source from papers that
            # merely cite it. This is a document prior, never invented evidence.
            if len(title_phrases[i]) >= 25 and title_phrases[i] in query_phrase:
                score *= 1.5
            # BM25 IDF is only ~0.29 in a one-passage library. A high fixed
            # threshold would incorrectly reject a precise match in that case.
            if score > 0.20:
                scores.append((score, i))
        ranking = sorted(scores, reverse=True)
        if encoder is not None and vectors:
            try:
                qv = self._unit(encoder([query])[0])
                if len(qv) != len(vectors[0]):
                    raise ValueError("查询与文档嵌入维度不一致")
                # Conservative hybrid: vector similarity reranks only lexical
                # candidates, avoiding an unconditional irrelevant top-k match.
                lexical = {index: 1 / (40 + rank) for rank, (_, index) in enumerate(ranking, 1)}
                semantic = sorted(((sum(a * b for a, b in zip(qv, vectors[i])), i)
                                   for _, i in ranking), reverse=True)
                for rank, (_, index) in enumerate(semantic, 1):
                    lexical[index] += 1 / (40 + rank)
                ranking = sorted(((score, index) for index, score in lexical.items()), reverse=True)
            except Exception as ex:
                self.embedding_warning = f"查询嵌入失败，本次使用 BM25：{ex}"
        if not ranking:
            return []
        # Suppress nearly identical overlapping chunks; retain different portions
        # of a page when the question requires several details.
        selected, seen_terms = [], []
        for score, index in ranking:
            if score < ranking[0][0] * 0.30:
                break
            candidate = current[index]
            # Preserve the surrounding explanation when the relevant sentence
            # falls at a chunk boundary. Neighbours never cross a PDF page or
            # slide, and the resulting citation ID identifies the actual text.
            neighbours = [j for j in range(max(0, index - 1), min(n, index + 2))
                          if current[j].source == candidate.source and current[j].location == candidate.location]
            pieces = [current[j].text for j in neighbours]
            context = pieces[0] if pieces else candidate.text
            for piece in pieces[1:]:
                overlap = next((length for length in range(min(220, len(context), len(piece)), 19, -1)
                                if context.endswith(piece[:length])), 0)
                context += piece[overlap:] if overlap else "\n" + piece
            context = _remove_contact_details(context)
            candidate = _passage(candidate.source, candidate.location, context)
            terms = set(tokenize(context))
            if any(current[index].source == p.source and current[index].location == p.location
                   and len(terms & prev) / max(len(terms | prev), 1) > 0.68
                   for p, prev in zip(selected, seen_terms)):
                continue
            selected.append(replace(candidate, score=round(score, 6)))
            seen_terms.append(terms)
            if len(selected) >= min(top_k, 10):
                break
        return selected
